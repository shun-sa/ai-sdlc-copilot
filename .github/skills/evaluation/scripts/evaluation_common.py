#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any

import yaml


def read_json(path: Path) -> Any:
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def read_yaml(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")
    with path.open("r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    if not isinstance(data, dict):
        raise ValueError(f"YAML root must be a mapping: {path}")
    return data


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def parse_datetime(value: Any) -> datetime:
    text = str(value or "").strip()
    if not text:
        raise ValueError("Datetime value is empty.")
    try:
        dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ValueError(f"Invalid ISO-8601 datetime: {text}") from exc
    if dt.tzinfo is None:
        raise ValueError(f"Datetime must include timezone offset: {text}")
    return dt


def pct(numerator: int | float, denominator: int | float) -> float | None:
    if denominator == 0:
        return None
    return round(float(numerator) / float(denominator) * 100.0, 2)


def as_upper(value: Any) -> str:
    return str(value or "").strip().upper()


def as_number(value: Any, *, default: float | None = None) -> float | None:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def load_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()

    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return [dict(row) for row in csv.DictReader(f)]

    if suffix == ".json":
        data = read_json(path)
        if isinstance(data, list):
            return [row for row in data if isinstance(row, dict)]
        if isinstance(data, dict):
            for key in ("rows", "cases", "items", "results"):
                rows = data.get(key)
                if isinstance(rows, list):
                    return [row for row in rows if isinstance(row, dict)]
        raise ValueError(f"JSON table must be an array or contain rows/cases/items/results: {path}")

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl is required to read .xlsx evaluation inputs."
            ) from exc

        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(v or "").strip() for v in values[0]]
        result: list[dict[str, Any]] = []
        for row in values[1:]:
            if all(v is None for v in row):
                continue
            result.append({
                headers[i]: row[i] if i < len(row) else None
                for i in range(len(headers))
                if headers[i]
            })
        return result

    raise ValueError(f"Unsupported table format: {path}")


def nested_get(data: dict[str, Any], dotted: str, default: Any = None) -> Any:
    current: Any = data
    for part in dotted.split("."):
        if not isinstance(current, dict) or part not in current:
            return default
        current = current[part]
    return current
